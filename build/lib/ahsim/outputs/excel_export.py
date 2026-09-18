"""Excel Report Engine (Spec Sections 59-66, 123).

26 mandatory sheets, professional formatting, traceability columns,
developer identity on cover and final report. All data comes from ONE
SimulationResult - never re-computed or copied from UI.
"""
from __future__ import annotations

import datetime
from typing import Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..app.metadata import (CONTACT, CREDIT_LINES, DEVELOPER_EN, DEVELOPER_FA,
                            MODEL_VERSIONS, PROJECT_NAME_DEFAULT, SOFTWARE_NAME,
                            VERSION)

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
WARN_FILL = PatternFill("solid", fgColor="FFC7CE")
IN_FILL = PatternFill("solid", fgColor="FFF2CC")
SRC_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHEETS = ["00_COVER", "01_PROJECT", "02_FEED", "03_MATERIAL", "04_MINERALOGY",
          "05_PSD", "06_EQUIPMENT", "07_FLOWSHEET", "08_STREAMS", "09_MASS_BALANCE",
          "10_FE_BALANCE", "11_RECOVERY", "12_PRE_PROCESSING", "13_SCREENING",
          "14_CRUSHING", "15_FINAL_MAGNETIC", "16_RECYCLE", "17_PRODUCTION",
          "18_SCENARIOS", "19_SENSITIVITY", "20_CALIBRATION", "21_ASSUMPTIONS",
          "22_DATA_SOURCES", "23_VALIDATION", "24_WARNINGS", "25_FINAL_REPORT"]


def _header(ws, row, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _table(ws, start_row, headers, rows, widths=None, number_formats=None):
    _header(ws, start_row, headers)
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if number_formats and c in number_formats and isinstance(v, (int, float)):
                cell.number_format = number_formats[c]
    if widths:
        for c, w in widths.items():
            col = c if isinstance(c, int) else c
            ws.column_dimensions[col].width = w
    ws.auto_filter.ref = (f"A{start_row}:"
                          f"{get_column_letter(len(headers))}{start_row + len(rows)}")


def export_workbook(result, path: str, project_name: str = PROJECT_NAME_DEFAULT,
                    plant_name: str = "Dry Magnetite Processing Plant",
                    scenario: str = "Scenario-001"):
    k = result.kpis
    st: Dict = result.streams
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        wb.create_sheet(name)

    # 00_COVER ------------------------------------------------------------
    ws = wb["00_COVER"]
    ws["B2"], ws["B3"] = SOFTWARE_NAME, "Industrial Mineral Processing Simulation and Engineering Software"
    ws["B2"].font = Font(bold=True, size=16)
    cover = [("Software Version", VERSION), ("Project Name", project_name),
             ("Plant Name", plant_name), ("Simulation Date", result.timestamp[:10]),
             ("Scenario", scenario), ("Simulation ID", result.simulation_id),
             ("Developer / Manufacturer", f"{DEVELOPER_EN} ({DEVELOPER_FA})"),
             ("Contact", CONTACT), ("Convergence Status", result.status)]
    _table(ws, 5, ["Field", "Value"], cover, widths={"A": 30, "B": 55})
    ws["B15"] = "MODEL RESULT - not a site-verified result until calibrated (Section 113)"
    ws["B15"].fill = WARN_FILL

    # 01_PROJECT ----------------------------------------------------------
    ws = wb["01_PROJECT"]
    _table(ws, 1, ["Item", "Value"], [
        ("Software", SOFTWARE_NAME), ("Version", VERSION),
        ("Model Versions", ", ".join(f"{a} {b}" for a, b in MODEL_VERSIONS.items())),
        ("Input Hash (SHA-256/16)", result.input_hash), ("Result Hash (SHA-256/16)", result.result_hash),
        ("Solver Status", result.solver.status), ("Iterations", result.solver.iterations),
        ("Developer", DEVELOPER_EN), ("Contact", CONTACT)])

    # 02_FEED -------------------------------------------------------------
    ws = wb["02_FEED"]
    cfg = result.config
    _table(ws, 1, ["Parameter", "Value", "Unit", "Source Type", "Basis"], [
        ("RAW FEED RATE", cfg.feed_rate_t_h, "t/h", "USER INPUT", "DRY"),
        ("Fe Grade", cfg.feed_fe_pct, "%", "USER INPUT", "DRY"),
        ("Moisture", cfg.moisture_pct, "%", "USER INPUT", "-"),
        ("Ore Density", cfg.ore_density_t_m3, "t/m3", "USER INPUT", "-"),
        ("Operating Hours/Day", cfg.operating_hours_day, "h", "USER INPUT", "-"),
        ("Operating Days/Month", cfg.operating_days_month, "d", "USER INPUT", "-"),
        ("Operating Days/Year", cfg.operating_days_year, "d", "USER INPUT", "-")],
        widths={"A": 28, "B": 14, "C": 10, "D": 22, "E": 10})

    # 03_MATERIAL ---------------------------------------------------------
    rows = [(c, "", "%", "USER INPUT", "") for c in
            ["Fe", "FeO", "Fe2O3", "SiO2", "Al2O3", "CaO", "MgO", "Mn", "P", "S", "Other"]]
    rows[0] = ("Fe", cfg.feed_fe_pct, "%", "USER INPUT", "DRY")
    ws = wb["03_MATERIAL"]
    _table(ws, 1, ["Assay Component", "Value", "Unit", "Source Type", "Basis"], rows,
           widths={"A": 20, "B": 12, "C": 8, "D": 22, "E": 8})

    # 04_MINERALOGY -------------------------------------------------------
    ws = wb["04_MINERALOGY"]
    _table(ws, 1, ["Mineral", "Fe Content % (stoichiometric)", "Magnetic", "Source"], [
        ("Magnetite", 72.36, "YES", "PUBLISHED RESEARCH"),
        ("Hematite", 69.94, "NO", "PUBLISHED RESEARCH"),
        ("Goethite", 62.86, "NO", "PUBLISHED RESEARCH"),
        ("Gangue / Quartz / Clay / Other", "", "NO", "DEFAULT VALUE")],
        widths={"A": 30, "B": 24, "C": 10, "D": 22})

    # 05_PSD --------------------------------------------------------------
    ws = wb["05_PSD"]
    psd = st["S-FEED"].psd
    _table(ws, 1, ["Size Fraction", "Size Lower (mm)", "Size Upper (mm)", "Mass %", "Fe %"],
           [(f.name, f.size_lower_mm, f.size_upper_mm, f.mass_pct, f.fe_pct) for f in psd.fractions],
           number_formats={4: "0.00", 5: "0.00"})

    # 06_EQUIPMENT --------------------------------------------------------
    ws = wb["06_EQUIPMENT"]
    eq_rows = [("GYRATORY", "PRIMARY GYRATORY", "CSS 200 mm", "USER INPUT", "CALIBRATION REQUIRED: capacity")]
    for s in [f"SEP-0{i}" for i in range(1, 7)]:
        eq_rows.append((s, "PRE-PROCESSING DRUM SEPARATOR", "R=1.5 m, D=3.0 m, L=3.0 m, shell 10 mm SS",
                        "PROJECT SPEC (Sections 18, 83)", "CALIBRATION REQUIRED: capacity, field profile"))
    for s in [f"SCR-0{i}" for i in range(1, 4)]:
        eq_rows.append((s, "SCREEN 3 mm", "2.4 x 6.0 m (user input)", "USER INPUT", ""))
    eq_rows += [("HYD-COARSE", "HYDROCONE CRUSHER (COARSE)", "CSS 35 mm (user input)", "USER INPUT", "UNKNOWN: power"),
                ("HYD-FINE", "HYDROCONE CRUSHER (FINE)", "CSS 5 mm (user input)", "USER INPUT", "UNKNOWN: power"),
                ("FIN-01", "FINAL DRUM SEPARATOR", "R=1.5 m, D=3.0 m", "PROJECT SPEC", "CALIBRATION REQUIRED"),
                ("FIN-02", "FINAL DRUM SEPARATOR", "R=1.5 m, D=3.0 m", "PROJECT SPEC", "CALIBRATION REQUIRED")]
    _table(ws, 1, ["Equipment ID", "Type", "Geometry / Setting", "Source", "Missing Data"],
           eq_rows, widths={"A": 14, "B": 32, "C": 34, "D": 28, "E": 34})

    # 07_FLOWSHEET --------------------------------------------------------
    ws = wb["07_FLOWSHEET"]
    seen, flow_rows = set(), []
    for sid, s in st.items():
        key = (s.source_equipment, s.destination_equipment)
        if key in seen:
            continue
        seen.add(key)
        flow_rows.append((sid, s.source_equipment, s.destination_equipment,
                          round(s.dry_mass_flow_t_h, 3)))
    _table(ws, 1, ["Stream ID", "Source", "Destination", "Dry Flow (t/h)"], flow_rows)

    # 08_STREAMS ----------------------------------------------------------
    ws = wb["08_STREAMS"]
    _table(ws, 1, ["Stream ID", "Wet t/h", "Dry t/h", "Moisture %", "Fe %", "Fe t/h", "Status"],
           [(sid, round(s.wet_mass_flow_t_h, 3), round(s.dry_mass_flow_t_h, 3),
             s.moisture_pct, round(s.fe_grade_pct, 3), round(s.fe_mass_flow_t_h, 3), s.status)
            for sid, s in st.items()], number_formats={2: "0.000", 3: "0.000", 6: "0.000"})

    # 09/10 balances --------------------------------------------------------
    ws = wb["09_MASS_BALANCE"]
    _table(ws, 1, ["Check", "Value"], [
        ("Total Feed (dry t/h)", k["feed_t_h"]), ("Total Product (dry t/h)", k["product_t_h"]),
        ("Total Tail (dry t/h)", k["tail_t_h"]),
        ("Mass Balance Error %", k["mass_balance_error_pct"]),
        ("PASS/FAIL", "PASS" if k["mass_balance_error_pct"] < 1.0 else "FAIL")])
    ws = wb["10_FE_BALANCE"]
    _table(ws, 1, ["Check", "Value"], [
        ("Fe In (t/h)", k["fe_in_t_h"]), ("Fe Product (t/h)", k["fe_product_t_h"]),
        ("Fe Tail (t/h)", k["fe_tail_t_h"]), ("Fe Balance Error %", k["fe_balance_error_pct"]),
        ("PASS/FAIL", "PASS" if k["fe_balance_error_pct"] < 1.0 else "FAIL")])

    # 11_RECOVERY -----------------------------------------------------------
    ws = wb["11_RECOVERY"]
    _table(ws, 1, ["Metric", "Value %"], [
        ("Mass Recovery", k["mass_recovery_pct"]), ("Fe Recovery", k["fe_recovery_pct"]),
        ("Mass Yield (product/feed)", k["mass_recovery_pct"]),
        ("Circulating Load %", k["circulating_load_pct"])])

    # 12-15 units -----------------------------------------------------------
    ws = wb["12_PRE_PROCESSING"]
    _table(ws, 1, ["Separator", "Feed t/h", "Magnetic t/h", "Tail t/h", "Magnetic Fe %", "Field G"],
           [(f"SEP-0{i}", round(st[f'S-PRE{i}-IN'].dry_mass_flow_t_h, 2),
             round(st[f'S-PRE{i}-MAG'].dry_mass_flow_t_h, 2),
             round(st[f'S-PRE{i}-TAIL'].dry_mass_flow_t_h, 2),
             round(st[f'S-PRE{i}-MAG'].fe_grade_pct, 2), 3000) for i in range(1, 7)])
    ws = wb["13_SCREENING"]
    _table(ws, 1, ["Screen", "Feed t/h", "Undersize t/h", "Oversize t/h", "Aperture mm"],
           [(f"SCR-0{i}", round(st[f'S-SCR{i}-IN'].dry_mass_flow_t_h, 2),
             round(st[f'S-SCR{i}-UND'].dry_mass_flow_t_h, 2),
             round(st[f'S-SCR{i}-OV'].dry_mass_flow_t_h, 2), 3.0) for i in range(1, 4)])
    ws = wb["14_CRUSHING"]
    _table(ws, 1, ["Crusher", "Feed t/h", "Product t/h", "CSS mm"], [
        ("HYD-COARSE", round(st["S-OV-COARSE"].dry_mass_flow_t_h, 2),
         round(st["S-COARSE-OUT"].dry_mass_flow_t_h, 2), 35.0),
        ("HYD-FINE", round(st["S-COARSE-OUT"].dry_mass_flow_t_h, 2),
         round(st["S-FINE-OUT"].dry_mass_flow_t_h, 2), 5.0)])
    ws = wb["15_FINAL_MAGNETIC"]
    _table(ws, 1, ["Separator", "Feed t/h", "Product t/h", "Tail t/h", "Product Fe %", "Field G"],
           [(f"FIN-0{i}", round(st[f'S-FIN{i}-IN'].dry_mass_flow_t_h, 2),
             round(st[f'S-PRODUCT-{i}'].dry_mass_flow_t_h, 2),
             round(st[f'S-TAIL-{i}'].dry_mass_flow_t_h, 2),
             round(st[f'S-PRODUCT-{i}'].fe_grade_pct, 2), 1400) for i in (1, 2)])

    # 16_RECYCLE ------------------------------------------------------------
    ws = wb["16_RECYCLE"]
    hist = result.solver.history
    _table(ws, 1, ["Iteration", "Recycle t/h", "Recycle Fe %", "|dm|", "|dFe|"],
           [(h["iteration"], round(h["recycle_t_h"], 3), round(h["recycle_fe_pct"], 4),
             h["dm"], h["dfe"]) for h in hist[-50:]])
    ws["F1"] = f"Status: {result.solver.status} | Iterations: {result.solver.iterations}"

    # 17_PRODUCTION ---------------------------------------------------------
    ws = wb["17_PRODUCTION"]
    _table(ws, 1, ["Metric", "Value"], [
        ("Product t/h", k["product_t_h"]), ("Product t/day", k["production_t_day"]),
        ("Product t/month", k["production_t_month"]), ("Product t/year", k["production_t_year"])],
        number_formats={2: "#,##0.00"})

    # 18-19 scenario/sensitivity --------------------------------------------
    ws = wb["18_SCENARIOS"]
    _table(ws, 1, ["Scenario", "Feed t/h", "Product t/h", "Product Fe %", "Fe Recovery %",
                   "Mass Recovery %", "Circulating Load t/h", "Status"],
           [(scenario, k["feed_t_h"], k["product_t_h"], round(k["product_fe_pct"], 2),
             round(k["fe_recovery_pct"], 2), round(k["mass_recovery_pct"], 2),
             round(k["circulating_load_t_h"], 1), result.status)])
    ws = wb["19_SENSITIVITY"]
    _table(ws, 1, ["Parameter", "Perturbation", "Fe Recovery %", "Product t/h", "Note"], [
        ("Feed Rate", "user scenarios 100-1000 t/h", "run scenario batch", "", "stress-tested, not capacity"),
        ("Magnetic Field", "3000/4000/5000 G pre", "partition curve dependent", "", "ENGINEERING CORRELATION"),
        ("Screen Efficiency", "0.7-1.0", "affects recycle", "", "USER INPUT")])

    # 20_CALIBRATION --------------------------------------------------------
    ws = wb["20_CALIBRATION"]
    _table(ws, 1, ["Test ID", "Date", "Feed Mass t", "Feed Fe %", "Field G", "RPM",
                   "Product Mass t", "Product Fe %", "Tail Mass t", "Tail Fe %",
                   "Observed Rec %", "Predicted Rec %", "Residual", "Rel. Error %"],
           [["CALIBRATION REQUIRED - enter plant test data; engine computes "
             "Observed/Predicted/Residual/Relative Error (Section 27)"]], widths={"A": 100})

    # 21_ASSUMPTIONS --------------------------------------------------------
    ws = wb["21_ASSUMPTIONS"]
    _table(ws, 1, ["#", "Assumption / Limitation"], [
        (1, "Month = 30 days, year = 365 days in unit conversion of t/month, t/year"),
        (2, "Primary gyratory modelled as crusher with 200 mm CSS (correlation model)"),
        (3, "Screen partition: logistic curve, sharpness from efficiency (Tier 5 correlation)"),
        (4, "Crusher breakage: Rosin-Rammler around CSS (Tier 5 correlation, needs testwork)"),
        (5, "Magnetic partition: size/liberation logistic (Tier 5) - MUST be calibrated (Sec 27-28)"),
        (6, "Crushing conserves Fe grade (no loss model)"),
        (7, "Final separator feed split equal between FIN-01/FIN-02"),
        (8, "Separator capacity, power, pole geometry = CALIBRATION REQUIRED / UNKNOWN (Sec 130)")],
        widths={"A": 5, "B": 110})

    # 22_DATA_SOURCES -------------------------------------------------------
    ws = wb["22_DATA_SOURCES"]
    _table(ws, 1, ["Parameter", "Value", "Source Type", "Tier", "Confidence"], [
        ("Drum geometry R/D/L", "1.5 / 3.0 / 3.0 m", "PROJECT SPEC", "-", "High"),
        ("Shell thickness", "10 mm Stainless Steel", "MANUFACTURER DATA", "2", "Medium"),
        ("Screen aperture", "3 mm", "USER INPUT", "6", "High"),
        ("Magnetic partition curve", "logistic model", "ENGINEERING CORRELATION", "5", "Low - CALIBRATION REQUIRED"),
        ("Crusher product PSD", "Rosin-Rammler", "ENGINEERING CORRELATION", "5", "Low - CALIBRATION REQUIRED"),
        ("Separator capacity", "UNKNOWN", "CALIBRATION REQUIRED", "-", "Unknown"),
        ("Equipment power", "UNKNOWN", "UNKNOWN", "-", "Unknown")])

    # 23_VALIDATION ---------------------------------------------------------
    ws = wb["23_VALIDATION"]
    _table(ws, 1, ["Validation", "Result"], [
        ("PSD Sum = 100%", "PASS" if abs(psd.total_mass_pct() - 100) < 0.01 else "FAIL"),
        ("Mass Balance < 1%", "PASS" if k["mass_balance_error_pct"] < 1 else "FAIL"),
        ("Fe Balance < 1%", "PASS" if k["fe_balance_error_pct"] < 1 else "FAIL"),
        ("Recovery in [0,100]", "PASS" if 0 <= k["fe_recovery_pct"] <= 100 else "FAIL"),
        ("Convergence", result.solver.status),
        ("Input Hash", result.input_hash), ("Result Hash", result.result_hash)])

    # 24_WARNINGS -----------------------------------------------------------
    ws = wb["24_WARNINGS"]
    warn = result.warnings or ["No warnings."]
    for i, w in enumerate(warn, 3):
        ws.cell(row=i, column=1, value=w).fill = WARN_FILL
    _header(ws, 1, ["Warnings"])

    # 25_FINAL_REPORT -------------------------------------------------------
    ws = wb["25_FINAL_REPORT"]
    _table(ws, 1, ["Metric", "Value", "Unit"], [
        ("Feed Rate", k["feed_t_h"], "t/h"), ("Feed Fe Grade", cfg.feed_fe_pct, "%"),
        ("Final Product Rate", k["product_t_h"], "t/h"), ("Final Product Grade", k["product_fe_pct"], "%"),
        ("Final Tail Rate", k["tail_t_h"], "t/h"), ("Final Tail Grade", k["tail_fe_pct"], "%"),
        ("Mass Recovery", k["mass_recovery_pct"], "%"), ("Fe Recovery", k["fe_recovery_pct"], "%"),
        ("Fe Input", k["fe_in_t_h"], "t/h"), ("Fe Product", k["fe_product_t_h"], "t/h"),
        ("Fe Tail", k["fe_tail_t_h"], "t/h"), ("Circulating Load", k["circulating_load_t_h"], "t/h"),
        ("Circulating Load", k["circulating_load_pct"], "%"),
        ("Screen Load", k["screen_load_t_h"], "t/h"), ("Crusher Load", k["crusher_load_t_h"], "t/h"),
        ("Daily Production", k["production_t_day"], "t/day"),
        ("Monthly Production", k["production_t_month"], "t/month"),
        ("Annual Production", k["production_t_year"], "t/year"),
        ("Mass Balance Error", k["mass_balance_error_pct"], "%"),
        ("Fe Balance Error", k["fe_balance_error_pct"], "%"),
        ("Convergence Status", result.status, ""),
        ("Simulation ID", result.simulation_id, ""),
        ("Developer / Manufacturer", f"{DEVELOPER_EN} ({DEVELOPER_FA})", ""),
        ("Contact", CONTACT, "")],
        widths={"A": 26, "B": 16, "C": 10}, number_formats={2: "#,##0.000"})
    for r in range(2, 26):
        c = ws.cell(row=r, column=2)
        if isinstance(c.value, str) and ("CONVERGED" in c.value or "WARNING" in c.value):
            if "NOT" in c.value or "WARNING" in c.value:
                c.fill = WARN_FILL

    # page setup + freeze panes
    for name in SHEETS:
        s = wb[name]
        s.freeze_panes = "A2" if name not in ("00_COVER",) else None
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_view.showGridLines = False

    wb.save(path)
    return path
